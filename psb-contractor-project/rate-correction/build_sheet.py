"""Build the before/after review workbook for one rate-correction batch.

Usage: python3 -I build_sheet.py <batch_dir> <ids_file> <out.xlsx> <batch_label>
"""
import html
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

batch_dir, ids_file, out, label = sys.argv[1:5]

RULES = [
    ("R1", "12.2% stated as the current combined Ontario small business rate", "Change to 11.2% and add \"(12.2% before July 1, 2026)\" once per page"),
    ("R2", "3.2% stated as the current Ontario small business rate", "Change to 2.2% and add \"(3.2% before July 1, 2026)\" once per page"),
    ("R3", "12.2% or 3.2% tied to 2025 or earlier, or to income before July 1, 2026", "Leave as is"),
    ("R4", "Dollar figure worked out from 12.2%", "Recalculate at 11.2%; Sharad confirms the figure"),
    ("R5", "Example for a December 31, 2026 year-end", "Use the blended 11.70%"),
    ("R6", "Calculator code with 0.122 or 0.032", "Replace with the day-prorated logic from the page 7 calculator"),
    ("R7", "Small business limit shown as $600,000", "Change to $500,000"),
]

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="0F2A4A")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
LEAVE_FILL = PatternFill("solid", fgColor="EEEEEE")
THIN = Side(style="thin", color="9AA5B1")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def clean(s):
    s = re.sub(r"<[^>]+>", "", s or "")
    s = html.unescape(s)
    return re.sub(r"\s+", " ", s).strip()


def style_table(ws, ncols, nrows, widths):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD, HEAD_FILL, WRAP, GRID
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r in range(2, nrows + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.alignment, cell.border = FONT, WRAP, GRID
    ws.freeze_panes = "A2"


ids = [l.split()[0] for l in open(ids_file) if l.strip()]
pages, rows = [], []
for pid in ids:
    path = os.path.join(batch_dir, f"{pid}.json")
    if not os.path.exists(path):
        pages.append({"id": int(pid), "url": "", "title": "", "page_note": "MISSING: extraction not done", "n": 0})
        continue
    d = json.load(open(path))
    n = 0
    for r in d.get("rows", []):
        n += 1
        rows.append((d, r))
    pages.append({"id": d["id"], "url": d.get("url", ""), "title": d.get("title", ""), "page_note": d.get("page_note", ""), "n": n})

wb = Workbook()

# Sheet 1: before / after
ws = wb.active
ws.title = "Before-After"
hdr = ["#", "Page URL", "Page ID", "Where", "Section", "Rule", "Before (current text)", "After (proposed)",
       "Note", "Sharad decision", "Sharad comment", "Raw before (for applying)", "Raw after (for applying)"]
ws.append(hdr)
for i, (d, r) in enumerate(rows, 1):
    ws.append([i, d.get("url", ""), d["id"], r.get("location", ""), r.get("section", ""), r.get("rule", ""),
               clean(r.get("before")), clean(r.get("after")), r.get("note", ""), "", "", r.get("before", ""), r.get("after", "")])
n = len(rows) + 1
style_table(ws, len(hdr), n, [5, 38, 9, 11, 16, 6, 60, 60, 30, 14, 30, 40, 40])
dv = DataValidation(type="list", formula1='"Approve,Change,Reject"', allow_blank=True)
ws.add_data_validation(dv)
for r in range(2, n + 1):
    dv.add(f"J{r}")
    ws[f"J{r}"].fill = INPUT_FILL
    ws[f"K{r}"].fill = INPUT_FILL
    if ws[f"F{r}"].value == "R3":
        for c in "ABCDEFGHI":
            ws[f"{c}{r}"].fill = LEAVE_FILL
ws.column_dimensions["L"].hidden = True
ws.column_dimensions["M"].hidden = True
ws.auto_filter.ref = f"A1:K{n}"

# Sheet 2: URLs in this batch
ws2 = wb.create_sheet("URLs")
ws2.append(["#", "Page URL", "Page ID", "Title", "Rows to review", "Rows to change", "Note"])
for i, p in enumerate(pages, 1):
    r = i + 1
    ws2.append([i, p["url"], p["id"], p["title"],
                f"=COUNTIF('Before-After'!$C:$C,C{r})",
                f"=COUNTIFS('Before-After'!$C:$C,C{r},'Before-After'!$F:$F,\"<>R3\")",
                p["page_note"]])
total = len(pages) + 2
ws2.cell(row=total, column=4, value="Total")
ws2.cell(row=total, column=5, value=f"=SUM(E2:E{total - 1})")
ws2.cell(row=total, column=6, value=f"=SUM(F2:F{total - 1})")
style_table(ws2, 7, total, [5, 55, 9, 45, 12, 12, 40])
for c in range(1, 8):
    ws2.cell(row=total, column=c).font = BOLD

# Sheet 3: rules and how to review
ws3 = wb.create_sheet("Rules")
ws3.append(["Rule", "Where the figure appears", "Action"])
for r in RULES:
    ws3.append(list(r))
style_table(ws3, 3, len(RULES) + 1, [8, 60, 60])
base = len(RULES) + 3
notes = [
    f"{label}. Source: Sharad's Existing Page Fixes and Rate Correction Rules (PSB package, Oct 3, 2026).",
    "How to review: in the Before-After tab, choose Approve, Change or Reject in the yellow 'Sharad decision' column. Use 'Sharad comment' for any change.",
    "Example: decision 'Change', comment 'Use 11.2% but drop the (12.2% before July 1, 2026) note here'.",
    "Grey rows are R3 (historical figures): proposed to stay as they are.",
    "Every visible FAQ change has a matching JSON-LD row so the schema stays identical to the page.",
    "Nothing on the website changes until Sharad approves the batch.",
]
for i, t in enumerate(notes):
    c = ws3.cell(row=base + i, column=1, value=t)
    c.font = FONT
    ws3.merge_cells(start_row=base + i, start_column=1, end_row=base + i, end_column=3)
    c.alignment = WRAP

wb.save(out)
print(len(pages), "pages,", len(rows), "rows")
